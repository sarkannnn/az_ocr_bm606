import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('results/Az_OCR_Akademik_Hesabat.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i:02d}: {p.text}")

print(f"\nTotal tables: {len(doc.tables)}")
for i, t in enumerate(doc.tables):
    print(f"\n--- Table {i+1} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
    for r_idx, r in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  Row {r_idx:02d}: {cells}")
