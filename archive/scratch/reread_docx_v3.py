import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('results/Az_OCR_Akademik_Hesabat_DUZELISLER_v3.docx')

print("="*80)
print("TITLE BLOCK")
print("="*80)
print("Title:", doc.paragraphs[0].text)
print("Meta:", doc.paragraphs[1].text)

print("\n" + "="*80)
print("ABSTRACT (XÜLASƏ)")
print("="*80)
print(doc.paragraphs[3].text)

print("\n" + "="*80)
print("TABLE 1: V1 vs V2 Augmentation Pipeline")
print("="*80)
t1 = doc.tables[0]
for i, r in enumerate(t1.rows):
    print(f"Row {i:02d}: | {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)} |")

print("\n" + "="*80)
print("TABLE 2: Training Metrics & Resource Consumption")
print("="*80)
t2 = doc.tables[1]
for i, r in enumerate(t2.rows):
    print(f"Row {i:02d}: | {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)} |")

print("\n" + "="*80)
print("TABLE 3: Synthetic Test Sets & Font Classification")
print("="*80)
t3 = doc.tables[2]
for i, r in enumerate(t3.rows):
    print(f"Row {i:02d}: | {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)} |")

print("\n" + "="*80)
print("TABLE 5: Real Document Test 2 Summary")
print("="*80)
t5 = doc.tables[4]
for i, r in enumerate(t5.rows):
    print(f"Row {i:02d}: | {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)} |")

print("\n" + "="*80)
print("TABLE 6: Real Document Direct Transcription Samples")
print("="*80)
t6 = doc.tables[5]
for i, r in enumerate(t6.rows):
    print(f"Row {i:02d}: | {' | '.join(c.text.strip().replace(chr(10), ' ') for c in r.cells)} |")

print("\n" + "="*80)
print("SECTION 1.2: Traditional OCR vs Deep Learning")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("1.2."):
        print(p.text)
    elif "Müasir Tesseract 4 və 5 mühərrikləri" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 1.3: Research Contributions")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("1.3."):
        print(p.text)
    elif "Azərbaycan Dili üçün İxtisaslaşmış CRNN Modeli" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 2.1: Text Corpus Formation & Leakage Prevention")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("2.1."):
        print(p.text)
    elif "Modelin təlimində Azərbaycan dilli Vikipediya" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 2.3: Bounding-Box Geometry & Measurement Bug Fix")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("2.3."):
        print(p.text)
    elif "make_sim_ladder.py skriptində söz kəsimlərinin" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 4.1: Standard Synthetic Mixed Tests")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("4.1."):
        print(p.text)
    elif "Standart sintetik testlərdə V2 modeli" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 4.2 Closing Sentence (with Tesseract L4 Word CER):")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Nəticə göstərir ki, V1 modeli real kamera"):
        print(p.text)

print("\n" + "="*80)
print("SECTION 5.1: Real Document Test 1")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("5.1."):
        print(p.text)
    elif "Birinci sınaqda kamera bucaq altında" in p.text:
        print(p.text)

print("\n" + "="*80)
print("SECTION 6.1: Groups A, B, C Description & Table 7")
print("="*80)
for p in doc.paragraphs:
    if "Audit 4 çərçivəsində şriftlər V2-nin təlim siyahısına nəzərən" in p.text:
        print(p.text)
    elif "Cədvəl 7 üzrə Mühüm Elmi Qeydlər" in p.text:
        print("\nNotes under Table 7:\n" + p.text)

print("\n" + "="*80)
print("SECTION 6.2: Rare Diacritics & Error Rows")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Audit 3 və Audit 4-ün çaşqınlıq matrisi"):
        print(p.text)
    elif p.text.startswith("Nadir Diakritik Xətalar"):
        print("\n" + p.text)
    elif p.text.startswith("Durğu işarələri və CTC boşluq xətaları"):
        print("\n" + p.text)

print("\n" + "="*80)
print("SECTION 7: Limitations (Items 1, 4, 7, 8)")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Tədqiqat Məhdudiyyətləri"):
        print(p.text)

print("\n" + "="*80)
print("REFERENCES")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith(("1. Shi", "2. Graves", "3. Smith", "4. Baek", "5.")):
        print(p.text)
