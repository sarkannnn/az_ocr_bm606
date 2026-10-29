import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('results/Az_OCR_Akademik_Hesabat_DUZELISLER_v2.docx')

print("==================================================")
print("1. TITLE BLOCK (P0 & P1):")
print("==================================================")
print("P0:", doc.paragraphs[0].text)
print("P1:", doc.paragraphs[1].text)

print("\n==================================================")
print("2. ABSTRACT (P3):")
print("==================================================")
print("P3:", doc.paragraphs[3].text)
print("P4 (checking deleted):", repr(doc.paragraphs[4].text))

print("\n==================================================")
print("3. TABLE 1 (Augmentations):")
print("==================================================")
t1 = doc.tables[0]
for r_idx, r in enumerate(t1.rows):
    print(f"Row {r_idx:02d}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")

print("\n==================================================")
print("4. TABLE 4 (Simulation Ladder):")
print("==================================================")
t4 = doc.tables[3]
for r_idx, r in enumerate(t4.rows):
    print(f"Row {r_idx:02d}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")

print("\n==================================================")
print("5. SEC 6.1 NEW MAIN TABLE (Groups A, B, C):")
print("==================================================")
# Find the new table (Table 7)
t7 = doc.tables[6]
for r_idx, r in enumerate(t7.rows):
    print(f"Row {r_idx:02d}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")

print("\n==================================================")
print("6. SEC 7 (Limitations & Future Work):")
print("==================================================")
for i, p in enumerate(doc.paragraphs):
    if "7. Tədqiqat Məhdudiyyətləri" in p.text or "Tədqiqat Məhdudiyyətləri (Reproducibility" in p.text:
        print(f"P{i:02d}: {p.text}\n")
    if "Dil Modeli İnteqrasiyası:" in p.text:
        print(f"P{i:02d}: {p.text}\n")

print("\n==================================================")
print("7. REFERENCE 1 & REFERENCE 5 CHECK:")
print("==================================================")
for i, p in enumerate(doc.paragraphs):
    if "1. Shi" in p.text:
        print(f"Ref 1 (P{i:02d}): {p.text}")
    if "5. Memmedov" in p.text or (len(p.text.strip()) > 0 and "Memmedov" in p.text):
        print(f"Ref 5 (P{i:02d}): {p.text}")
