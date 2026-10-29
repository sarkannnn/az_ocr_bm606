import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('results/Az_OCR_Akademik_Hesabat_DUZELISLER_v4.docx')

print("="*80)
print("1. ABSTRACT (XÜLASƏ)")
print("="*80)
print(doc.paragraphs[3].text)

print("\n" + "="*80)
print("2. TABLE 2 (ALL ROWS & COLS)")
print("="*80)
t2 = doc.tables[1]
for i, r in enumerate(t2.rows):
    print(f"Row {i:02d}: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in r.cells) + " |")

print("\n" + "="*80)
print("3. TABLE 4 (ALL ROWS & COLS)")
print("="*80)
t4 = doc.tables[3]
for i, r in enumerate(t4.rows):
    print(f"Row {i:02d}: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in r.cells) + " |")

print("\n" + "="*80)
print("4. TABLE 7 (ALL CELLS)")
print("="*80)
t7 = doc.tables[6]
for i, r in enumerate(t7.rows):
    print(f"Row {i:02d}: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in r.cells) + " |")

print("\n" + "="*80)
print("5. TABLE 8 (ALL CELLS)")
print("="*80)
t8 = doc.tables[7]
for i, r in enumerate(t8.rows):
    print(f"Row {i:02d}: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in r.cells) + " |")

print("\n" + "="*80)
print("6. SECTION 7 (ALL LIMITATIONS)")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Tədqiqat Məhdudiyyətləri"):
        print(p.text)

print("\n" + "="*80)
print("7. VERIFICATION OF CHANGED SECTIONS (Sec 1.2, 1.3, 2.1, 4.1, 4.2, 5.1, 6.1, 6.2, Ref 3, Table 1 Row 4)")
print("="*80)
for i, p in enumerate(doc.paragraphs):
    if any(p.text.startswith(prefix) for prefix in [
        "1.2.", "1.3.", "Müasir Tesseract", "Azərbaycan Dili üçün İxtisaslaşmış",
        "2.1.", "Modelin təlimində Azərbaycan", "4.1.", "Standart sintetik",
        "Nəticə göstərir ki, V1 modeli", "5.1.", "Birinci sınaqda",
        "6.1.", "Audit 4 çərçivəsində şriftlər", "Cədvəl 7 üzrə Mühüm",
        "6.2.", "Audit 3 və Audit 4-ün çaşqınlıq", "Nadir Diakritik Xətalar",
        "3. Tesseract OCR"
    ]):
        print(f"P{i:02d}: {p.text}\n")

print("\nTABLE 1 ROW 4 (Salt & Pepper):")
t1 = doc.tables[0]
print(f"Row 04: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in t1.rows[4].cells) + " |")
