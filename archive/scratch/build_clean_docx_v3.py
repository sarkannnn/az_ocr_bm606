import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

SRC_DOCX = Path("results/Az_OCR_Akademik_Hesabat_DUZELISLER_v2.docx")
DST_DOCX = Path("results/Az_OCR_Akademik_Hesabat_DUZELISLER_v3.docx")

shutil.copy2(SRC_DOCX, DST_DOCX)
doc = docx.Document(DST_DOCX)
print(f"Copied {SRC_DOCX} -> {DST_DOCX}")

def set_run_font(run, name="Calibri", size=10.0, color=None, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def format_cell(cell, text, bold=False, italic=False, size=8.5, bg_color=None):
    cell.text = text
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    for r in p.runs:
        set_run_font(r, size=size, bold=bold, italic=italic)
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

# ---------------------------------------------------------------------
# 1. Title Block (P0, P1)
# ---------------------------------------------------------------------
doc.paragraphs[0].text = "Azərbaycan Dili üçün Dərin Öyrənmə Əsaslı Optik Simvol Tanıma (CRNN+CTC) Sistemi və Real Dünya Sənəd Təsvirlərində Qiymətləndirilməsi"
for r in doc.paragraphs[0].runs:
    set_run_font(r, name="Calibri", size=16, bold=True, color=RGBColor(0x1B, 0x36, 0x5D))

doc.paragraphs[1].text = "Müəllif: [Müəllif Adı] / BM 606 Tədqiqat Qrupu | Layihə Kodu: BM 606 | Tarix: 29 Oktyabr 2026\nElmi Hesabat və Metodoloji Audit Sənədi"
for r in doc.paragraphs[1].runs:
    set_run_font(r, name="Calibri", size=9.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))

# ---------------------------------------------------------------------
# 2. Abstract (P3)
# ---------------------------------------------------------------------
doc.paragraphs[3].text = (
    "Bu tədqiqat işində Azərbaycan dilinin xüsusi qrafemlərini (ə, ı, ö, ğ, ü, ş, ç) ehtiva edən çap sənədlərinin "
    "rəqəmsallaşdırılması üçün Konvolusion və İkiistiqamətli Təkrarlanan Neyron Şəbəkələrinə (CRNN) və Bağlantıçı Zaman "
    "Təsnifatına (CTC) əsaslanan dərin öyrənmə sistemi layihələndirilmiş, təlim edilmiş və empirik audit olunmuşdur. "
    "Kamera təhrifləri ilə zənginləşdirilmiş V2 modeli simulyasiya şəraitində ilkin V1 modelini kəskin üstələyir (V2 >> V1). "
    "Bununla belə, klaster-bootstrap təhlili (95% CI) göstərir ki, V2 modeli yalnız təlim şrift ailələri üzrə Tesseract 5 "
    "ilə bərabər səviyyədədir; tamamilə yeni şrift ailələrində isə L4 sözdə Tesseract öndədir (+1.55% [+0.59, +2.94]); "
    "digər 3 müqayisədə fərq əhəmiyyətli deyil. Real mobil telefon çəkilişlərində Tesseract 5 sətir səviyyəsində aydın şəkildə "
    "öndədir: Foto 1-də Tesseract CER 21.92% vs V2 28.87%; Foto 2-də Tesseract CER 0.24% (19/21 sətir tam qüsursuz) vs V2 5.02%. "
    "Real sınaq eyni bir fiziki sənədin iki fərqli fotosundan ibarət olduğundan nəticələr ilkin xarakter daşıyır; bundan əlavə, "
    "Foto 1 kəsimləri avtomatik icra edildiyi üçün qismən sürüşmüşdür (sətir #4 etalon uyğunsuzluğu) və Foto 2 sətir CER göstəricisi "
    "kəsim keyfiyyəti baxımından tam təsdiqlənməmişdir (21 sətir kəsiminin 17-də yuxarı/aşağı kənarda mürəkkəb toxunması mövcuddur). "
    "Sintetik qarışıq testlərdə CER-in 0.50%-dən 0.11%-ə və 1.64%-dən 0.30%-ə enməsi test dəstinin yarısının Consolas ailə sızması "
    "hesabına baş vermişdir; digər yarısı olan Tahoma şriftində isə V2 heç bir irəliləyiş göstərməmişdir (0.08% vs 0.09% söz; 0.24% vs 0.26% sətir)."
)
for r in doc.paragraphs[3].runs:
    set_run_font(r, name="Calibri", size=10, italic=True)

# ---------------------------------------------------------------------
# 3. Body Paragraphs Updates
# ---------------------------------------------------------------------
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    
    # Sec 2.3 Heading & Body
    if txt.startswith("2.3. Bounding-Box Həndəsəsi"):
        p.text = "2.3. Bounding-Box Həndəsəsi və Ölçmə Xətasının Aradan Qaldırılması"
        for r in p.runs:
            set_run_font(r, size=12, bold=True, color=RGBColor(0x1B, 0x36, 0x5D))
            
    elif "tight ink bounding box" in txt or "84.32%-dən 0.76%-ə" in txt:
        p.text = (
            "Audit 3 zamanı aşkar edilmişdir ki, make_sim_ladder.py skriptində söz kəsimlərinin koordinatlarının hesablanmasında "
            "proqram xətası (measurement bug) mövcud olmuşdur: şaquli font metrikası ilə (ascent + descent = 80px) hesablanan qutular "
            "qonşu sətir və sözlərin mürəkkəbini kəsimin içinə salırdı. Bu ölçmə xətası düzəldilərək sıx mürəkkəb qutusu (tight ink bounding box) "
            "tətbiq edildikdə, SansSerifCollection şriftində ölçülən CER göstəricisi 84.32%-dən 0.76%-ə düşmüşdür; bu dəyişiklik "
            "modelin təkmilləşməsi deyil, ölçmə alətindəki proqram qüsurunun aradan qaldırılmasıdır."
        )
        for r in p.runs:
            set_run_font(r, size=10)

    # Sec 4.2 closing sentence: add Tesseract L4 word CER (0.37%)
    elif txt.startswith("Nəticə göstərir ki, V1 modeli real kamera şəraitində (L4) 35.77% CER ilə yararsız vəziyyətə"):
        p.text = (
            "Nəticə göstərir ki, V1 modeli real kamera şəraitində (L4) 35.77% CER ilə yararsız vəziyyətə düşdüyü halda, "
            "V2 modeli sözdə 1.35% CER qoruyub saxlayır (müqayisə üçün: Tesseract 5 bu pillədə 0.37% CER nümayiş etdirir). "
            "Sətir səviyyəsində isə L4-də V2 modeli 0.11% CER ilə Tesseract 5-dən (0.33%) daha dayanıqlı nəticə göstərir."
        )
        for r in p.runs:
            set_run_font(r, size=10)

    # Sec 6.1 explanation and font list for Groups A, B, C
    elif "Audit 4 çərçivəsində şriftlər" in txt and "3 qrupa ayrılaraq" in txt:
        p.text = (
            "Audit 4 çərçivəsində şriftlər V2-nin təlim siyahısına nəzərən 3 qrupa ayrılaraq tam klaster-bootstrap metodologiyası (95% CI) "
            "ilə qiymətləndirilmişdir:\n"
            "• Qrup A (Seen, 4 şrift: arial.ttf, calibri.ttf, segoeui.ttf, times.ttf): V2 təlimində tam iştirak edən baza şriftlər.\n"
            "• Qrup B (Near-Relative, 5 şrift: calibrib.ttf, CascadiaMono.ttf, pala.ttf, seguisb.ttf, verdanab.ttf): Təlim şriftləri ilə birbaşa ailə qohumluğu olan (leakage) şriftlər.\n"
            "• Qrup C (New Families, 7 şrift: cambria.ttc, micross.ttf, Moba_Fanta.ttf, Moba_Fira.ttf, Moba_Font.ttf, SansSerifCollection.ttf, tahomabd.ttf): Təlimdə heç bir forması iştirak etməyən tamamilə yeni şrift ailələri.\n\n"
            "Qeyd: Şrift qrupları V2 modelinin 14 təlim şriftinə nəzərən müəyyən edilmişdir; V1 üçün Qrup A şriftləri də söz səviyyəsində faktiki olaraq yenidir, çünki V1-in söz təlim korpusunda rəqəm və durğu işarələri 0% olmuşdur (bu səbəbdən V1 Qrup A L0 söz CER göstəricisi 7.36% təşkil edir).\n\n"
            "Cədvəl 7. Şrift qrupları (A, B, C) üzrə L0 və L4 səviyyələrində V1, V2 və Tesseract 5 CER müqayisəsi və klaster-bootstrap 95% etibarlılıq intervalları (Qruplar V2-nin təlim siyahısına nəzərən müəyyən edilmişdir; V1 üçün Qrup A da faktiki olaraq yenidir):"
        )
        for r in p.runs:
            set_run_font(r, size=10)

    # Sec 6.1 note under table
    elif "Cədvəl 7 üzrə Mühüm Elmi Qeydlər" in txt:
        p.text = (
            "Cədvəl 7 üzrə Mühüm Elmi Qeydlər (Şrift qrupları V2 təlim siyahısına görə təyin edilmişdir; V1 üçün Qrup A da faktiki olaraq yenidir):\n"
            "1. Qrup A Word L0 üzrə V2 CER (0.81%, 20/2472 hərf) ilə Qrup C Word L4 üzrə Tesseract CER (0.81%, 34/4194 hərf) göstəricilərinin eyni olması sırf rəqəmsal təsadüfdür; hər iki göstərici yoxlanmış birbaşa saylardan qaynaqlanır.\n"
            "2. Qrup C üzrə 4 müqayisədən yalnız 1-də (Word L4) fərq intervalı sıfırı əhatə etmir (+1.55% [+0.59%, +2.94%]); bu pillədə Tesseract V2 modelini birmənalı üstələyir. Digər 3 halda modellər statistik bərabərdir.\n"
            "3. Qrup A (k=4) və Qrup B (k=5) klaster sayları az olduğu üçün bu qrupların intervalları yüksək dispersiyaya malikdir. Moba_Fira monospace şriftidir və təlimdə CascadiaCode/consolab iştirak etmişdir."
        )
        for r in p.runs:
            set_run_font(r, size=9.0, italic=True)

    # Sec 6.2 rare cases & error analysis - remove file:line tags
    elif txt.startswith("Audit 3 və Audit 4-ün çaşqınlıq matrisi analizi göstərir ki, V2 modelində 'ə' -> 'e' xətası 0.00%"):
        p.text = (
            "Audit 3 və Audit 4-ün çaşqınlıq matrisi analizi göstərir ki, V2 modelində 'ə' -> 'e' xətası 0.00% təşkil edir (0 / 394 instansiya). "
            "Dotless-i (ı/i/İ) üzrə V2 modelində qeydə alınan cəmi 106 xətanın analizi (Audit 4):\n"
            "• Birbaşa 'ı' <-> 'i' əvəzləməsi cəmi 5 halda baş vermişdir (4 ədəd 'i->ı', 1 ədəd 'ı->i').\n"
            "• Əsas xətalar 'i' -> 'ü' (58 hal) və 'ı' -> 'u' (20 hal) qarışıqlığından ibarətdir.\n"
            "• Bu xətaların 91.5%-i (97 / 106) cəmi 3 şriftdə cəmlənmişdir: Moba_Fanta (53 hal / 50.0%), Moba_Font (30 hal / 28.3%) və SansSerifCollection (14 hal / 13.2%). "
            "Əvvəlki layihə qaralamasında qeyd olunan 'ı->i 14.2% -> 0.8%' və 'ə->e 8.7% -> 0.4%' iddiaları təsdiqlənməmişdir (mənbəsizdir)."
        )
        for r in p.runs:
            set_run_font(r, size=10)

    elif txt.startswith("Nadir Diakritik Xətalar"):
        p.text = (
            "Nadir Diakritik Xətalar (Audit 4 nəticə fayllarından faktiki qeydə alınmış sətirlər):\n"
            "• ö -> o xətaları (o_ö_errors.csv):\n"
            "  - ref: 'Köpək' -> hyp: 'Kopək' | Model: V2, Qrup: C, Şrift: SansSerifCollection.ttf, Səviyyə: L4, Rejim: word\n"
            "  - ref: 'köynək-baju' -> hyp: 'koynək.baju' | Model: V2, Qrup: C, Şrift: SansSerifCollection.ttf, Səviyyə: L4, Rejim: word\n"
            "• o -> ə xətası (o_ö_errors.csv):\n"
            "  - ref: 'Goranboy,' -> hyp: 'Gəranboy,' | Model: V2, Qrup: C, Şrift: cambria.ttc, Səviyyə: L4, Rejim: word\n"
            "• c -> e xətaları (c_ç_errors.csv):\n"
            "  - ref: 'dcmokratik' -> hyp: 'demokratik' | Model: V2, Qrup: C, Şrift: cambria.ttc, Səviyyə: L4, Rejim: word\n"
            "  - ref: 'çevrilmiş, dcmokratik ruhlu və' -> hyp: 'çevrilmiş, demokratik ruhlu və' | Model: V2, Qrup: C, Şrift: cambria.ttc, Səviyyə: L4, Rejim: line\n"
            "• C -> c registri (c_ç_errors.csv):\n"
            "  - ref: 'C.' -> hyp: 'c.' | Model: V2, Qrup: B, Şrift: CascadiaMono.ttf, Səviyyə: L0/L4, Rejim: word\n"
            "• ç -> s əvəzləməsi (c_ç_errors.csv):\n"
            "  - ref: 'qaçan' -> hyp: 'qasan' | Model: V2, Qrup: B, Şrift: verdanab.ttf, Səviyyə: L4, Rejim: word\n"
            "• G -> C qarışması (c_ç_errors.csv):\n"
            "  - ref: 'Gürcüstan baş hakiminin Qarabağda' -> hyp: 'Cürcüstan baş hakiminin Qarabağda' | Model: V2, Qrup: B, Şrift: pala.ttf, Səviyyə: L0/L4, Rejim: line"
        )
        for r in p.runs:
            set_run_font(r, size=9.5)

    elif txt.startswith("Durğu işarələri və CTC boşluq xətaları"):
        p.text = (
            "Durğu işarələri və CTC boşluq xətaları: V2 sətir modelində sətir sonundakı tək nöqtə simvolu bəzən '..' kimi dekodlaşdırılır (məsələn, 'hazırlanması..', 'qurulması..').\n\n"
            "Sətir WER Boşluq Təhlili (Groups B & C, eyni 1076 söz üzrə xəta növləri dekompozisiyası):\n"
            "V2 modelinin sətir söz xətalarının böyük hissəsi (L0-da 79.3%, L4-də 65.3%) xalis boşluq xətalarından (həm sözlərin birləşməsi, həm də söz daxilinə səhvən boşluq daxil edilməsi) qaynaqlanır. Tesseract-da isə boşluq xətası demək olar ki, yoxdur (0.0% – 5.9%).\n\n"
            "Cədvəl 8. Sətir WER xəta növlərinin dekompozisiyası (Groups B & C, 1076 söz üzrə V2 və Tesseract 5 müqayisəsi):"
        )
        for r in p.runs:
            set_run_font(r, size=10)

    # Sec 7: Full Limitations
    elif txt.startswith("Tədqiqat Məhdudiyyətləri (Reproducibility & Scope Limitations):"):
        p.text = (
            "Tədqiqat Məhdudiyyətləri (Reproducibility & Scope Limitations):\n"
            "1. Render skriptinin provenansı: V2 verilənlər bazasına əlavə olunan 15,000 söz və 6,000 sətir mütləq yolu "
            "C:\\Users\\serxan.m\\.gemini\\antigravity-ide\\brain\\912dbc3a-1632-49b2-a686-a0d9534b585d\\scratch\\build_v2_datasets.py "
            "olan skript vasitəsilə generasiya edilmişdir (Dəyişdirilmə tarixi: 2026-10-23 11:29:19, SHA256: dc1fd149d6a39d4f60630fc9159ab8384ae2de294f5a0bd7ac3abab354083352). "
            "Skript generate_data.py daxilindəki make_split funksiyasını və TRDG GeneratorFromStrings mühərrikini çağırmışdır.\n"
            "2. Val dəstinin məhdudiyyəti: Val dəsti yalnız V1-in 8 təlim şriftini ehtiva edir (synth_word_v2/val və synth_line_v2/val); "
            "modelin ən yaxşı checkpoint-i yalnız görünən şriftlər üzrə seçilmişdir (unseen generalizasiya üzrə deyil).\n"
            "3. Təlim itkisinin qeyri-sabitliyi: V2 sətir modelində təlimin son epoxlarında kəskin val-loss sıçrayışları "
            "(epox 9-da 0.72, epox 13-də 0.65, epox 15-də 0.48) baş vermiş və model 10-cu epoxdan sonra erkən dayandırılmışdır (patience=5).\n"
            "4. Foto 1 və Foto 2 kəsimləri, etalon mənşəyi: Foto 1 kəsimləri avtomatik icra edildiyi üçün qismən sürüşmüşdür (sətir #4 etalon uyğunsuzluğu). "
            "Foto 2 kəsim qutuları isə Tesseract image_to_data çıxarışının təhlilindən sonra test_real_doc2.py faylına əllə daxil edilmişdir (müstəqil deyil). "
            "Etalon mətnlər sənəd fotosundan əllə transkripsiya edilmişdir (müstəqil cüt yoxlama sənədləşdirilmədiyi üçün status: 'təsdiq gözləyir'). "
            "21 sətir kəsiminin 17-də (81.0%) 2px marja daxilində şaquli kənarda davam edən mürəkkəb mövcuddur; 4-6px marja və ağ maskalama sınağı isə "
            "Otsu binarizasiyasını pozaraq (V2 9.08%, Tess 29.27%) nəticə verməmişdir, buna görə sətir CER göstəricisi kəsim keyfiyyəti baxımından tam təsdiqlənməmişdir.\n"
            "5. Söz kəsimi və qonşu müdaxiləsi: 30 söz kəsiminin 18-də (60.0%) qonşu sözlərin mürəkkəb müdaxiləsi (neighbour intrusion >15px) "
            "mövcuddur; kəsilmiş söz benchmarkı etibarsız sayılır.\n"
            "6. Real sınaq həcminin məhdudluğu: Yalnız eyni bir fiziki sənədin 2 fotoşəkli sınaqdan keçirilmişdir; 20-30 foto və ya 100+ sətirlik "
            "genişmiqyaslı real sənəd testi hələ aparılmamışdır.\n"
            "7. Təhrifli və qəliz şriftlərdə Tesseract-ın üstünlüyü: Tesseract 5 həm real sənəd sətirlərində (0.24% vs 5.02%), həm nərdivan L4 "
            "sözlərində (0.37% vs 1.35%), həm də qəliz əl/dekorativ şriftlərdə (Inkfree L0 söz: Tesseract 14.20% vs V2 30.69%) layihə modelindən üstündür. "
            "Tahoma sətirlərində isə fərq nominaldır və statistik əhəmiyyətli deyil (Tesseract 0.22% vs V2 0.26%, n=500).\n"
            "8. Test kəsimlərində generator xətası: make_sim_ladder.py skriptində PIL ImageFont qutularının tərtibində kəsim xətası səbəbindən "
            "sözlərin 2.3%-də (24/1049 kəsim, məs. 'və' sözündə 11.8% mürəkkəb itkisi) mürəkkəb kəsilməsi qeydə alınmışdır. Xüsusi hərflərin seçmə həcmi kiçikdir (ç üçün n=23, ş üçün n=56).\n"
            "9. Rəqəm və durğu işarələri: Yalnız V2-yə əlavə olunan 15,000 sözün 20%-i (3,000 söz) rəqəm və durğu işarələrini ehtiva edir; "
            "ilkin V1 təlim korpusunda rəqəm və durğu ümumiyyətlə olmamışdır."
        )
        for r in p.runs:
            set_run_font(r, size=9.5)

# ---------------------------------------------------------------------
# 4. Clean Table 1
# ---------------------------------------------------------------------
t1 = doc.tables[0]
while len(t1.rows) < 13:
    t1.add_row()

t1_headers = ["Parametr / Təhrif Növü", "V1 Modeli (Baza)", "V2 Modeli (Fotorealistik Kamera Borusu)"]
for c_idx, h in enumerate(t1_headers):
    format_cell(t1.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF")

t1_rows = [
    ["Fırlanma Bucağı (Rotation)", "± 5° incə fırlanma (p=0.5)", "± 5° incə fırlanma (p=0.5)"],
    ["Perspektiv Əyilmə (Perspective)", "± 2% en, ± 8% hündürlük (p=0.3)", "± 2% en, ± 8% hündürlük (p=0.3)"],
    ["Qauss Bulanığı (Gaussian Blur)", "k ∈ {3, 5} (p=0.3)", "σ ∈ [0.5, 2.0] (p=0.6)"],
    ["Duz-İstiot Küyü (Salt & Pepper)", "0.5% duz + 0.5% istiot (1.0% piksel, p=0.3)", "0.5% duz + 0.5% istiot (1.0% piksel, p=0.2)"],
    ["Hərəkət Bulanığı (Motion Blur)", "Yoxdur", "kernel 3–7 (p=0.15)"],
    ["Sensor/Qauss Küyü (Sensor Noise)", "Yoxdur", "σ ∈ [3.0, 12.0] (p=0.6)"],
    ["Rezolyusiya İtkisi (Downscale)", "Yoxdur", "faktor 0.5× – 0.9× (p=0.4)"],
    ["JPEG Sıxılma Artefaktı", "Yoxdur", "Q ∈ [40, 90] (p=0.4)"],
    ["Parlaqlıq və Kontrast", "Yoxdur", "α ∈ [0.6, 1.1], β ∈ [-40, 20] (p=0.6)"],
    ["İşıq Qradiyenti (Illumination)", "Yoxdur", "hamar qradiyent 0.7 – 1.0 (p=0.5)"],
    ["Yumşaq Kölgə Zolağı (Dark Band)", "Yoxdur", "qradiyent daxilində p=0.5 (ümumi p=0.25), dərinlik 12%–28%, en 0.2–0.5"],
    ["Kontrast Dartılması", "Yoxdur", "min-max normalizasiyası (p=0.5)"],
]
for r_idx, row_data in enumerate(t1_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t1.rows[r_idx].cells[c_idx], val)

# ---------------------------------------------------------------------
# 5. Clean Table 2
# ---------------------------------------------------------------------
t2 = doc.tables[1]
t2_headers = ["Model", "Ümumi Epox", "Ən Yaxşı Epox", "Val Loss", "Val CER (%)*", "Təlim Vaxtı", "Checkpoint Ölçüsü"]
for c_idx, h in enumerate(t2_headers):
    format_cell(t2.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF")

t2_rows = [
    ["V2 Word", "18 (Early Stop, pat=5)", "Epox 13", "0.0063", "0.24% (Seen fonts)*", "212.3 dəq ümumi (154.1 dəq @ best)", "98.7 MB (xalis: 32.9 MB)"],
    ["V2 Line", "15 (Early Stop, pat=5)", "Epox 10", "0.0125", "0.28% (Seen fonts)*", "280.4 dəq ümumi (186.9 dəq @ best)", "98.7 MB (xalis: 32.9 MB)"],
]
for r_idx, row_data in enumerate(t2_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t2.rows[r_idx].cells[c_idx], val)

# ---------------------------------------------------------------------
# 6. Clean Table 3
# ---------------------------------------------------------------------
t3 = doc.tables[2]
t3_headers = ["Dataset / Test Çoxluğu", "Dəqiq Şrift Tərkibi və Təsnifatı", "V1 Modeli CER", "V2 Modeli CER", "Tesseract 5 CER", "V1 -> V2 Dəyişimi"]
for c_idx, h in enumerate(t3_headers):
    format_cell(t3.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF")

t3_rows = [
    ["Söz Testi (word_synth)", "1,001 tahoma (New) + 999 consola (Leakage)", "0.50%", "0.11%", "0.70%", "CER 4.55x azaldı (Qazanc yalnız Consolas sızmasındandır, Tahoma dəyişməyib)"],
    ["Sətir Testi (line_synth)", "501 tahoma (New) + 499 consola (Leakage)", "1.64%", "0.30%", "0.44%", "CER 5.47x azaldı (Qazanc yalnız Consolas sızmasındandır, Tahoma dəyişməyib)"],
    ["Söz, Tahoma (diagw_tahoma)", "1,000 tahoma.ttf — NEW FAMILY", "0.08%", "0.09%", "0.50%", "Dəyişmədi (+0.01% dalğalanma)"],
    ["Sətir, Tahoma (diag_tahoma)", "500 tahoma.ttf — NEW FAMILY", "0.24%", "0.26%", "0.22%", "Dəyişmədi (+0.02%, Tesseract öndədir)"],
    ["Söz, Consolas (diagw_consola)", "1,000 consola.ttf — FAMILY LEAKAGE (consolab)", "1.02%", "0.07%", "0.86%", "CER 14.57x azaldı (Ailə sızması)"],
    ["Sətir, Consolas (diag_consola)", "500 consola.ttf — FAMILY LEAKAGE (consolab)", "3.48%", "0.27%", "0.47%", "CER 12.89x azaldı (Ailə sızması)"],
]
for r_idx, row_data in enumerate(t3_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t3.rows[r_idx].cells[c_idx], val)

# ---------------------------------------------------------------------
# 7. Clean Table 4
# ---------------------------------------------------------------------
t4 = doc.tables[3]
t4_headers = ["Pillə", "Təhrif Növü", "V1 Word CER", "V2 Word CER", "Tess Word CER", "V1 Line CER", "V2 Line CER", "Tess Line CER"]
for c_idx, h in enumerate(t4_headers):
    format_cell(t4.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF", size=8.0)

t4_rows = [
    ["L0", "Təmiz render (Sıfır təhrif)", "7.71%", "0.20%", "0.41%", "1.87%", "0.11%", "0.22%"],
    ["L1", "İşıqlanma qradiyenti və kölgə", "8.85%", "0.20%", "0.45%", "4.10%", "0.11%", "0.22%"],
    ["L2", "Bulanıqlıq və sensor küyü", "25.00%", "0.41%", "0.41%", "13.84%", "0.07%", "0.29%"],
    ["L3", "Ölçü kiçiltmə (0.6x) və JPEG", "31.36%", "0.53%", "0.41%", "19.14%", "0.18%", "0.22%"],
    ["L4", "Tam kamera simulyasiyası", "35.77%", "1.35%", "0.37%", "12.96%", "0.11%", "0.33%"],
    ["L4n", "L4 + Kontrast dartılması", "24.43%", "1.14%", "0.37%", "7.43%", "0.07%", "0.33%"],
]
for r_idx, row_data in enumerate(t4_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t4.rows[r_idx].cells[c_idx], val, size=8.0)

# ---------------------------------------------------------------------
# 8. Clean Table 5
# ---------------------------------------------------------------------
t5 = doc.tables[4]
t5_headers = ["Kateqoriya", "Model", "CER (%)", "WER (%) [Standart]", "100% Dəqiq Tanınan Sayı / Cəmi"]
for c_idx, h in enumerate(t5_headers):
    format_cell(t5.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF")

t5_rows = [
    ["Sətirlər (21 sətir, 837 simvol)", "V1 Line", "48.51%", "98.88% (88/89)", "0 / 21 (0.0%)"],
    ["", "V2 Line", "5.02%", "22.47% (20/89)", "7 / 21 tam raw (9/21 strip)"],
    ["", "Tesseract 5", "0.24%", "2.25% (2/89)", "19 / 21 tam qüsursuz (90.5%)"],
    ["Sözlər (30 söz, 304 simvol)** [ETİBARSIZ]", "V1 Word", "90.13%", "100.00% (30/30)", "0 / 30 (0.0%) [Neighbour intrusion: 18/30]"],
    ["", "V2 Word", "46.71%", "90.00% (27/30)", "3 / 30 (10.0%) [Neighbour intrusion: 18/30]"],
    ["", "Tesseract 5", "50.99%", "146.67% (44/30)†", "4 / 30 (13.3%) [Neighbour intrusion: 18/30]"],
]
for r_idx, row_data in enumerate(t5_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t5.rows[r_idx].cells[c_idx], val)

# ---------------------------------------------------------------------
# 9. Clean Table 6
# ---------------------------------------------------------------------
t6 = doc.tables[5]
t6_headers = ["Orijinal Mətn (Ground Truth)", "V1 Line Modeli", "V2 Line Modeli", "Tesseract 5 (aze)"]
for c_idx, h in enumerate(t6_headers):
    format_cell(t6.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF")

t6_rows = [
    ["TƏQDİMAT (Başlıq, x=300..485)", "jirəggitifTijj", " TəqDiMAT ı", "TƏQDİMAT (100% Dəqiq)"],
    ["hazırlanması.", "ihagiianmasiji", "hazırlanması.. (Nöqtə təkrarlanması)", "hazırlanması."],
    ["Elektron jurnal idarəetmə sisteminin təkmilləşdirilməsi", "RiBlaktonijürmaiidaraatmaisistamin...", "Elektron jurnal idarəetmə sisteminin təkmilləşdirilməsi (100%)", "Elektron lurnal idarəetmə sisteminin..."],
    ["Nəşr və xidmət haqları barədə məlumatların ayrıca təqdim edilməsi.", "iHaşrijaixidməthağlaribarəğamalüin...", "Nəşr və xldmət haqları barədə məlumatların... ('i'->'l')", "Nəşr və xidmət haqları barədə məlumatların..."],
    ["Veb-saytın dizaynının və istifadəçi interfeysinin yenilənməsi", "y6bsaytinidizaynininipaistifadəği...", "y6b-saytın dizaynının və İstifadəçi interfeysinin yenilənməsi", "Vöb-saytın dizaynının və istifadəçi interfeysinin yenilənməsi"],
]
for r_idx, row_data in enumerate(t6_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t6.rows[r_idx].cells[c_idx], val)

# ---------------------------------------------------------------------
# 10. Clean Table 7 (Groups A, B, C)
# ---------------------------------------------------------------------
t7 = doc.tables[6]
t7_headers = ["Qrup", "Rejim və Səviyyə", "Nümunə (n)", "Simvol", "V1 CER (%)", "V2 CER [Cluster 95% CI]", "Tess CER [Cluster 95% CI]", "Δ (V2 - Tess) [Cluster 95% CI]", "0 CI-yə Daxildir?"]
for c_idx, h in enumerate(t7_headers):
    format_cell(t7.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF", size=7.5)

t7_rows = [
    ["Qrup A", "Word L0", "378", "2472", "7.36%", "0.81% [0.32%, 1.57%]", "2.39% [1.39%, 3.83%]", "-1.58% [-2.68%, -0.76%]", "XEYR (V2 öndədir)"],
    ["Qrup A", "Line L0", "95", "2777", "0.29%", "0.11% [0.00%, 0.30%]", "0.29% [0.00%, 0.65%]", "-0.18% [-0.57%, +0.11%]", "BƏLİ (Bərabər)"],
    ["Qrup A", "Word L4", "378", "2472", "45.27%", "1.33% [0.58%, 2.20%]", "1.70% [0.87%, 2.84%]", "-0.36% [-1.25%, +0.50%]", "BƏLİ (Bərabər)"],
    ["Qrup A", "Line L4", "95", "2777", "13.83%", "1.15% [0.18%, 3.00%]", "1.62% [0.11%, 4.09%]", "-0.47% [-1.91%, +0.44%]", "BƏLİ (Bərabər)"],
    ["Qrup B", "Word L0", "436", "3012", "9.56%", "0.46% [0.07%, 0.97%]", "0.96% [0.30%, 1.89%]", "-0.50% [-1.28%, +0.09%]", "BƏLİ (Bərabər)"],
    ["Qrup B", "Line L0", "114", "3368", "2.35%", "0.30% [0.00%, 0.92%]", "0.12% [0.00%, 0.36%]", "+0.18% [-0.31%, +0.88%]", "BƏLİ (Bərabər)"],
    ["Qrup B", "Word L4", "436", "3012", "41.33%", "0.73% [0.27%, 1.31%]", "1.10% [0.39%, 2.08%]", "-0.37% [-1.26%, +0.50%]", "BƏLİ (Bərabər)"],
    ["Qrup B", "Line L4", "114", "3368", "16.12%", "0.50% [0.06%, 1.17%]", "0.18% [0.00%, 0.44%]", "+0.33% [-0.18%, +1.06%]", "BƏLİ (Bərabər)"],
    ["Qrup C", "Word L0", "609", "4194", "9.73%", "1.62% [0.22%, 3.44%]", "1.00% [0.36%, 1.78%]", "+0.62% [-0.35%, +2.08%]", "BƏLİ (Bərabər)"],
    ["Qrup C", "Line L0", "158", "4673", "6.03%", "0.71% [0.08%, 1.68%]", "0.71% [0.19%, 1.46%]", "+0.00% [-0.74%, +0.76%]", "BƏLİ (Tam bərabər)"],
    ["Qrup C", "Word L4", "609", "4194", "41.49%", "2.36% [1.23%, 3.95%]", "0.81% [0.37%, 1.31%]", "+1.55% [+0.59%, +2.94%]", "XEYR (Tess öndədir)"],
    ["Qrup C", "Line L4", "158", "4673", "20.93%", "1.28% [0.16%, 2.75%]", "1.46% [0.37%, 3.72%]", "-0.17% [-2.60%, +1.53%]", "BƏLİ (Bərabər)"],
]
for r_idx, row_data in enumerate(t7_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t7.rows[r_idx].cells[c_idx], val, size=7.5)

# ---------------------------------------------------------------------
# 11. Clean Table 8 (WER Decomposition)
# ---------------------------------------------------------------------
t8 = doc.tables[7]
t8_headers = ["Səviyyə və Model", "Ümumi Söz (n)", "Xalis Boşluq Xətası", "Tək Hərf Xətası", "Çox-Hərf Xətası", "Səhv Söz Sayı", "Raw WER (%)", "Boşluq Xətası Payı (%)"]
for c_idx, h in enumerate(t8_headers):
    format_cell(t8.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF", size=8.0)

t8_rows = [
    ["L0 — V2 Line", "1076", "23", "3", "3", "29", "2.70%", "79.3%"],
    ["L0 — Tesseract 5", "1076", "0", "15", "21", "36", "3.35%", "0.0%"],
    ["L4 — V2 Line", "1076", "32", "10", "7", "49", "4.55%", "65.3%"],
    ["L4 — Tesseract 5", "1076", "2", "11", "21", "34", "3.16%", "5.9%"],
]
for r_idx, row_data in enumerate(t8_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t8.rows[r_idx].cells[c_idx], val, size=8.0)

doc.save(DST_DOCX)
print(f"Successfully finalized {DST_DOCX}")
