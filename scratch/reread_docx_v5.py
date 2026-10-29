import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('results/Az_OCR_Akademik_Hesabat_DUZELISLER_v5.docx')

print("="*80)
print("1. ABSTRACT (XÜLASƏ)")
print("="*80)
print(doc.paragraphs[3].text)

print("\n" + "="*80)
print("2. TABLE 1 ROW 4 (SALT & PEPPER PROSE)")
print("="*80)
t1 = doc.tables[0]
print(" | ".join(c.text.strip().replace(chr(10), " ") for c in t1.rows[4].cells))

print("\n" + "="*80)
print("3. TABLE 8 (ALL CELLS RECOMPUTED)")
print("="*80)
t8 = doc.tables[7]
for i, r in enumerate(t8.rows):
    print(f"Row {i:02d}: | " + " | ".join(c.text.strip().replace(chr(10), " ") for c in r.cells) + " |")

print("\n" + "="*80)
print("4. SEC 6.1 (XALİS HƏRF CER MƏXRƏCİ İLƏ)")
print("="*80)
for p in doc.paragraphs:
    if "Audit 4 çərçivəsində şriftlər V2-nin təlim siyahısına" in p.text:
        print(p.text)
    elif "Cədvəl 7 üzrə Mühüm Elmi Qeydlər" in p.text:
        print("\nQeyd 1:")
        print(p.text[:220])

print("\n" + "="*80)
print("5. SEC 6.2 (NADİR DİAKRİTİK XƏTALAR, S<->C VƏ DCMOKRATİK İZAHI)")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Nadir Diakritik Xətalar"):
        print(p.text)

print("\n" + "="*80)
print("6. SEC 7 (MƏHDUDİYYƏT 1 VƏ 4)")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("Tədqiqat Məhdudiyyətləri"):
        lines = p.text.split("\n")
        for l in lines:
            if l.startswith(("1.", "4.")):
                print(l)

print("\n" + "="*80)
print("7. İSTİNAD 3 (MÜRACİƏT TARİXİ İLƏ)")
print("="*80)
for p in doc.paragraphs:
    if p.text.startswith("3. Tesseract OCR"):
        print(p.text)
