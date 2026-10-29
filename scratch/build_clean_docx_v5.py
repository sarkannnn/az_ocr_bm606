import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = Path("results/Az_OCR_Akademik_Hesabat_DUZELISLER_v5.docx")
doc = docx.Document(DOCX_PATH)

def set_run_font(run, name="Calibri", size=10.0, color=None, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def format_cell(cell, text, bold=False, italic=False, size=8.5, bg_color=None):
    cell.text = text
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    for r in p.runs:
        set_run_font(r, size=size, bold=bold, italic=italic)
    if bg_color:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

# ---------------------------------------------------------------------
# 1. Update Section 6.1 (Source of letter denominator: 2300 letters)
# ---------------------------------------------------------------------
for p in doc.paragraphs:
    txt = p.text.strip()
    if "Audit 4 çərçivəsində şriftlər V2-nin təlim siyahısına" in txt:
        p.text = (
            "Audit 4 çərçivəsində şriftlər V2-nin təlim siyahısına nəzərən 3 qrupa ayrılaraq tam klaster-bootstrap metodologiyası (95% CI) "
            "ilə qiymətləndirilmişdir:\n"
            "• Qrup A (Seen, 4 şrift: arial.ttf, calibri.ttf, segoeui.ttf, times.ttf): Həm V1, həm V2 təlimində tam iştirak edən baza şriftlər.\n"
            "• Qrup B (Near-Relative, 5 şrift: calibrib.ttf, CascadiaMono.ttf, pala.ttf, seguisb.ttf, verdanab.ttf): Təlim şriftləri ilə birbaşa ailə qohumluğu olan (leakage) şriftlər.\n"
            "• Qrup C (New Families, 7 şrift: cambria.ttc, micross.ttf, Moba_Fanta.ttf, Moba_Fira.ttf, Moba_Font.ttf, SansSerifCollection.ttf, tahomabd.ttf): Təlimdə heç bir forması iştirak etməyən tamamilə yeni şrift ailələri.\n\n"
            "Qeyd: Qrup A şriftləri (`arial`, `calibri`, `segoeui`, `times`) V1-in öz təlim şriftləridir. Buna baxmayaraq, V1-in Qrup A Word L0 səviyyəsində "
            "7.36% CER göstərməsinin dəqiq səbəbi xəta bölgüsündən aydın olur: 378 test sözündəki 2,472 simvolun 2,300-ü əlifba hərfi, 172-si isə rəqəm və durğu işarəsidir "
            "(data/sim_audit4/seen/L0/word/test/labels.tsv). V1 söz təlim korpusunda rəqəm və durğu 0% olduğu üçün 172 rəqəm və durğu işarəsinin hamısında "
            "(172/172, 100%) xəta baş vermişdir (110 əvəzetmə və 62 silinmə). Əlifba hərfləri üzrə isə 2,300 hərfdə cəmi 10 xəta qeydə alınmışdır "
            "(məxrəc: 10/2300 hərf = 0.43% xalis hərf CER; ümumi 2,472 simvola nisbətdə 10/2472 = 0.40%). "
            "Eyni şriftlərin sətir modelində (Line L0) isə V1 durğu və rəqəmlərlə təlim edildiyi üçün CER cəmi 0.29% təşkil edir.\n\n"
            "Cədvəl 7. Şrift qrupları (A, B, C) üzrə L0 və L4 səviyyələrində V1, V2 və Tesseract 5 CER müqayisəsi və klaster-bootstrap 95% etibarlılıq intervalları (Qruplar V2-nin təlim siyahısına nəzərən müəyyən edilmişdir):"
        )
        for r in p.runs:
            set_run_font(r, size=10)

    elif txt.startswith("Cədvəl 7 üzrə Mühüm Elmi Qeydlər"):
        p.text = (
            "Cədvəl 7 üzrə Mühüm Elmi Qeydlər (Şrift qrupları V2 təlim siyahısına görə təyin edilmişdir):\n"
            "1. Qrup A şriftləri V1-in öz təlim şriftləridir; V1-in L0 söz xətalarının 94.5%-i (172/182) korpusda 0% olan rəqəm və durğu işarələridir (10/2300 hərf = 0.43% xalis hərf CER; ümumi 10/2472 = 0.40%).\n"
            "2. Qrup A Word L0 üzrə V2 CER (0.81%, 20/2472 hərf) ilə Qrup C Word L4 üzrə Tesseract CER (0.81%, 34/4194 hərf) göstəricilərinin eyni olması sırf rəqəmsal təsadüfdür; hər iki göstərici yoxlanmış birbaşa saylardan qaynaqlanır.\n"
            "3. Qrup C üzrə 4 müqayisədən yalnız 1-də (Word L4) fərq intervalı sıfırı əhatə etmir (+1.55% [+0.59%, +2.94%]); bu pillədə Tesseract V2 modelini birmənalı üstələyir. Digər 3 halda modellər statistik bərabərdir.\n"
            "4. Qrup A (k=4) və Qrup B (k=5) klaster sayları az olduğu üçün bu qrupların intervalları yüksək dispersiyaya malikdir. Moba_Fira monospace şriftidir və təlimdə CascadiaCode/consolab iştirak etmişdir."
        )
        for r in p.runs:
            set_run_font(r, size=9.0, italic=True)

    # ---------------------------------------------------------------------
    # 2. Update Section 6.2 (Delete Tesseract s<->c sentence, add dcmokratik delta impact)
    # ---------------------------------------------------------------------
    elif txt.startswith("Nadir Diakritik Xətalar"):
        p.text = (
            "Nadir Diakritik Xətalar (Audit 4 çaşqınlıq cədvəlindən dəqiq saylar və sətirlər):\n"
            "• ö -> o xətaları (ö üçün cəmi n=32 instansiya, V2-də 2 xəta [6.25%]; Tesseract-da 0 xəta):\n"
            "  - ref: 'Köpək' -> hyp: 'Kopək' | Model: V2, Qrup: C, Şrift: SansSerifCollection.ttf, Səviyyə: L4, Rejim: word\n"
            "  - ref: 'köynək-baju' -> hyp: 'koynək.baju' | Model: V2, Qrup: C, Şrift: SansSerifCollection.ttf, Səviyyə: L4, Rejim: word\n"
            "  (Tesseract 5 hər iki sözü 'Köpək' və 'köynək-baju' olaraq 100% dəqiq oxumuşdur).\n"
            "• o -> ə xətası (V2-də 1 xəta; Tesseract-da 0 xəta):\n"
            "  - ref: 'Goranboy,' -> hyp: 'Gəranboy,' | Model: V2, Qrup: C, Şrift: cambria.ttc, Səviyyə: L4, Rejim: word\n"
            "  (Tesseract 5 'Goranboy,' sözünü dəqiq tanımışdır).\n"
            "• ç -> s xətası (ç üçün cəmi n=23 instansiya, V2-də 1 xəta [4.35%]; Tesseract-da 0 xəta):\n"
            "  - ref: 'qaçan' -> hyp: 'qasan' | Model: V2, Qrup: B, Şrift: verdanab.ttf, Səviyyə: L4, Rejim: word\n"
            "  (Tesseract 5 'qaçan' sözünü tam dəqiq oxumuşdur).\n"
            "• c -> e halları (Etalon mətndə hərf səhvi / reference typo):\n"
            "  - ref: 'dcmokratik' -> hyp: 'demokratik' | Həm V2, həm Tesseract 5 (cambria.ttc, word və line)\n"
            "  (Qeyd: Etalon korpusda 'demokratik' sözü səhvən 'dcmokratik' daxil edilmişdir. Bu etalon xətası həm V2, həm Tesseract 5 üçün Qrup C L4 CER göstəricilərinə bərabər şəkildə daxil olduğundan, modellər arasındakı fərqə [Δ = V2 - Tess] heç bir təsir göstərmir. Həmçinin bu iki hal c_ç_errors.csv cədvəlində mövcuddur [c->e], lakin xalis 'ç' sayına [n=23] daxil deyil, çünki sözdə 'ç' deyil, 'c' hərfi yazılmışdır; 'o' hərfində xəta olmadığı üçün o_ö_errors.csv cədvəlinə və n=32 'ö' sayına da daxil deyil).\n"
            "• Digər Registr və Simvol Fərqləri:\n"
            "  - ref: 'C.' -> hyp: 'c.' | Həm V2, həm Tesseract (CascadiaMono.ttf, L0/L4, word)\n"
            "  - ref: 'Gürcüstan baş hakiminin Qarabağda' -> hyp: 'Cürcüstan baş hakiminin Qarabağda' | V2 (pala.ttf, L0/L4, line; G->C qarışması)"
        )
        for r in p.runs:
            set_run_font(r, size=9.5)

# ---------------------------------------------------------------------
# 3. Update Table 8: Column 7 Header -> "Səhv söz payı (%)"
# ---------------------------------------------------------------------
t8 = doc.tables[7]
t8_headers = ["Səviyyə və Model", "Ümumi Söz (n)", "Xalis Boşluq Xətası", "Tək Hərf Xətası", "Çox-Hərf Xətası", "Səhv Söz Sayı", "Səhv söz payı (%)", "Boşluq Xətası Payı (%)"]
for c_idx, h in enumerate(t8_headers):
    format_cell(t8.rows[0].cells[c_idx], h, bold=True, bg_color="EAECEF", size=8.0)

t8_rows = [
    ["L0 — V2 Line", "1076", "46", "8", "4", "58", "5.39%", "79.3%"],
    ["L0 — Tesseract 5", "1076", "0", "22", "7", "29", "2.70%", "0.0%"],
    ["L4 — V2 Line", "1076", "64", "26", "8", "98", "9.11%", "65.3%"],
    ["L4 — Tesseract 5", "1076", "2", "23", "9", "34", "3.16%", "5.9%"],
]
for r_idx, row_data in enumerate(t8_rows, start=1):
    for c_idx, val in enumerate(row_data):
        format_cell(t8.rows[r_idx].cells[c_idx], val, size=8.0)

doc.save(DOCX_PATH)
print(f"Successfully finalized {DOCX_PATH}")
